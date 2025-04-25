import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import csv, os, re, requests, shutil, subprocess
from datetime import date
from tkcalendar import DateEntry
import sv_ttk
from PIL import Image, ImageTk, ImageOps
from io import BytesIO
import openpyxl
import win32com.client
import traceback

# Ścieżki – ustaw według swoich potrzeb
BASE_IMPORT_PATH = r"S:\_KOLEKCJE_\IMPORT"
TEMPLATE_FILE = r"S:\Graficy Specyfikacje\Techniczne\szablony_specyfikacje różne\BIN\SPECYFIKACJA_IMPORT nie otwierac tutaj tylko kopiowac.indd"
EXCEL_FILE = r"G:\_Projekty\projekty.xlsx"
EXCEL_ICON_PATH = r"S:\Graficy Specyfikacje\Techniczne\_PROGRAMY\program files\worknavigate\icons\icons8-excel-32.png"

# ----------------------------------------
# Funkcje pomocnicze do menu kontekstowego
# ----------------------------------------
def add_text_context_menu(text_widget):
    menu = tk.Menu(text_widget, tearoff=0)
    menu.add_command(label="Wklej", command=lambda: text_widget.event_generate("<<Paste>>"))
    def show_menu(event):
        menu.tk_popup(event.x_root, event.y_root)
    text_widget.bind("<Button-3>", show_menu)

def add_entry_context_menu(entry):
    menu = tk.Menu(entry, tearoff=0)
    def paste():
        try:
            entry.insert(tk.END, entry.clipboard_get())
        except tk.TclError:
            pass
    menu.add_command(label="Wklej", command=paste)
    entry.bind("<Button-3>", lambda event: menu.tk_popup(event.x_root, event.y_root))

# ----------------------------------------
# ExcelTab – zakładka prezentująca dane z arkusza Excel
# ----------------------------------------
class ExcelTab(tk.Frame):
    def __init__(self, master, **kwargs):
        super().__init__(master, **kwargs)
        self.data = []
        self.filtered_data = []
        self.create_widgets()
        self.load_data()
    
    def create_widgets(self):
        search_frame = tk.Frame(self)
        search_frame.pack(fill="x", padx=10, pady=5)
        tk.Label(search_frame, text="Szukaj:").pack(side="left", padx=5)
        self.search_var = tk.StringVar()
        self.search_entry = ttk.Entry(search_frame, textvariable=self.search_var)
        self.search_entry.pack(side="left", padx=5, fill="x", expand=True)
        self.search_entry.bind("<KeyRelease>", lambda event: self.filter_data())
        refresh_btn = ttk.Button(search_frame, text="Odśwież", command=self.load_data)
        refresh_btn.pack(side="left", padx=5)
        
        self.scroll_frame = CustomScrolledFrame(self)
        self.scroll_frame.pack(fill="both", expand=True, padx=10, pady=5)
        self.rows_container = self.scroll_frame.inner
        
    def load_data(self):
        self.data = []
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] and row[5]:
                    self.data.append((str(row[0]), str(row[5])))
            wb.close()
        except Exception as e:
            messagebox.showerror("Błąd", f"Nie udało się odczytać arkusza Excel:\n{EXCEL_FILE}\n\n{e}")
        self.filtered_data = self.data.copy()
        self.populate_rows()
    
    def populate_rows(self):
        for child in self.rows_container.winfo_children():
            child.destroy()
        for item in self.filtered_data:
            self.add_row(item)
    
    def add_row(self, data_item):
        row_frame = tk.Frame(self.rows_container, bd=1, relief="solid", padx=5, pady=2)
        row_frame.pack(fill="x", padx=2, pady=2)
        name_label = tk.Label(row_frame, text=data_item[0], anchor="w")
        name_label.pack(side="left", fill="x", expand=True)
        action_btn = ttk.Button(row_frame, text="Idź do", command=lambda p=data_item[1]: self.open_path(p))
        action_btn.pack(side="right", padx=5)
    
    def open_path(self, path):
        fixed_path = path.replace("/", "\\")
        if os.path.exists(fixed_path):
            subprocess.Popen(["explorer", fixed_path])
        else:
            messagebox.showwarning("Uwaga", f"Ścieżka nie istnieje:\n{fixed_path}")
    
    def filter_data(self):
        query = self.search_var.get().lower()
        self.filtered_data = [item for item in self.data if query in item[0].lower()]
        self.populate_rows()

# ----------------------------------------
# FolderCreationFrame – tworzenie folderu oraz obsługa linku "przepisy prania nowy.ai"
# ----------------------------------------
class FolderCreationFrame(tk.Frame):
    def __init__(self, master, **kwargs):
        super().__init__(master, **kwargs)
        self.create_widgets()
    
    def create_widgets(self):
        wrapper = ttk.LabelFrame(self, text="Tworzenie folderu")
        wrapper.pack(fill="both", expand=True, padx=10, pady=10)
        
        season_frame = tk.Frame(wrapper)
        season_frame.pack(padx=10, pady=5, fill="x")
        tk.Label(season_frame, text="Wybierz folder sezonowy:").pack(side="left", padx=5)
        self.season_cb = ttk.Combobox(season_frame, state="readonly")
        self.season_cb.pack(side="left", padx=5, fill="x", expand=True)
        
        employee_frame = tk.Frame(wrapper)
        employee_frame.pack(padx=10, pady=5, fill="x")
        tk.Label(employee_frame, text="Wybierz folder pracownika:").pack(side="left", padx=5)
        self.employee_cb = ttk.Combobox(employee_frame, state="readonly")
        self.employee_cb.pack(side="left", padx=5, fill="x", expand=True)
        
        self.update_season_list()
        self.season_cb.bind("<<ComboboxSelected>>", self.update_employee_list)
        
        new_folder_frame = tk.Frame(wrapper)
        new_folder_frame.pack(padx=10, pady=5, fill="x")
        tk.Label(new_folder_frame, text="Nazwa nowego folderu:").pack(side="left", padx=5)
        self.new_folder_entry = ttk.Entry(new_folder_frame)
        self.new_folder_entry.pack(side="left", padx=5, fill="x", expand=True)
        
        create_btn = ttk.Button(wrapper, text="Utwórz folder", command=self.create_project_folder)
        create_btn.pack(padx=10, pady=10)
    
    def update_season_list(self):
        try:
            all_items = os.listdir(BASE_IMPORT_PATH)
            seasons = [item for item in all_items if os.path.isdir(os.path.join(BASE_IMPORT_PATH, item)) and item.upper().startswith("SEZON_")]
            self.season_cb['values'] = seasons
            if seasons:
                self.season_cb.current(0)
                if hasattr(self, 'employee_cb'):
                    self.update_employee_list()
        except Exception as e:
            messagebox.showerror("Błąd", f"Nie można odczytać folderów z:\n{BASE_IMPORT_PATH}\n\n{e}")
    
    def update_employee_list(self, event=None):
        season = self.season_cb.get()
        if season:
            season_path = os.path.join(BASE_IMPORT_PATH, season)
            try:
                employees = [item for item in os.listdir(season_path) if os.path.isdir(os.path.join(season_path, item))]
                self.employee_cb['values'] = employees
                if employees:
                    self.employee_cb.current(0)
            except Exception as e:
                messagebox.showerror("Błąd", f"Nie można odczytać folderów w:\n{season_path}\n\n{e}")
    
    def create_project_folder(self):
        season = self.season_cb.get()
        employee = self.employee_cb.get()
        new_folder_name = self.new_folder_entry.get().strip()
        
        if not (season and employee and new_folder_name):
            messagebox.showwarning("Uwaga", "Wybierz folder sezonowy, folder pracownika oraz wpisz nazwę nowego folderu.")
            return
        
        target_path = os.path.join(BASE_IMPORT_PATH, season, employee, new_folder_name)
        try:
            os.makedirs(target_path, exist_ok=True)
            os.makedirs(os.path.join(target_path, "WIZ"), exist_ok=True)
            os.makedirs(os.path.join(target_path, "CODE"), exist_ok=True)
            if os.path.exists(TEMPLATE_FILE):
                _, ext = os.path.splitext(TEMPLATE_FILE)
                new_template_name = f"{new_folder_name} tech file{ext}"
                target_file = os.path.join(target_path, new_template_name)
                shutil.copy2(TEMPLATE_FILE, target_file)
                # Po skopiowaniu szablonu wykonujemy operację kopiowania oraz relinkowania TYLKO dla linku "przepisy prania nowy.ai"
                self.reinit_przepisy_link(target_file, target_path)
            else:
                messagebox.showwarning("Uwaga", f"Plik szablonu nie istnieje:\n{TEMPLATE_FILE}")
            messagebox.showinfo("Sukces", f"Folder projektu utworzony:\n{target_path}")
            subprocess.Popen(["explorer", target_path])
        except Exception as e:
            messagebox.showerror("Błąd", f"Nie udało się utworzyć folderu:\n{target_path}\n\n{e}")
    
    def reinit_przepisy_link(self, indd_file, target_path):
        """
        Uruchamia ExtendScript (JavaScript dla InDesign), który:
        - Otwiera dokument INDD (ścieżka przekazana jako indd_file)
        - Szuka wszystkich łączy zawierających "przepisy prania nowy.ai"
        - Kopiuje oryginalny plik łącza do folderu LINKS w target_path (tylko raz)
        - Relinkuje wszystkie znalezione łącza do skopiowanego pliku
        - Zapisuje i zamyka dokument
        """
        try:
            indesign = win32com.client.Dispatch("InDesign.Application")
        except Exception as e:
            messagebox.showerror("Błąd", "Nie udało się połączyć z InDesign.\n" + str(e))
            return

        # Upewnij się, że folder LINKS istnieje w target_path
        links_folder = os.path.join(target_path, "LINKS")
        if not os.path.exists(links_folder):
            try:
                os.makedirs(links_folder)
            except Exception as e:
                messagebox.showerror("Błąd", f"Nie udało się utworzyć folderu LINKS:\n{links_folder}\n{e}")
                return

        # Przygotuj ścieżkę folderu LINKS w formacie odpowiednim dla ExtendScript (używamy ukośników "/")
        dest_folder_js = links_folder.replace("\\", "/")
        # Upewnij się, że ścieżka do dokumentu INDD jest w formacie z ukośnikami "/"
        indd_file_js = indd_file.replace("\\", "/")
        
        # Skrypt ExtendScript:
        # - Otwiera dokument INDD
        # - Iteruje przez wszystkie linki, szuka pasujących do targetLinkName
        # - Kopiuje plik (tylko przy pierwszym wystąpieniu) i relinkuje wszystkie
        # - Zapisuje i zamyka dokument
        js = """
    var doc = app.open(File("%s"));
    var targetLinkName = "przepisy prania nowy.ai";
    var destFolder = new Folder("%s");
    if (!destFolder.exists) {
        destFolder.create();
    }
    var newFile = null;
    for (var i = 0; i < doc.links.length; i++) {
        var link = doc.links[i];
        if (link.name.toLowerCase().indexOf(targetLinkName.toLowerCase()) !== -1) {
            if (newFile == null) {
                var srcFile = new File(link.filePath);
                if (srcFile.exists) {
                    newFile = new File(destFolder.fsName + "/" + srcFile.name);
                    srcFile.copy(newFile.fsName);
                }
            }
            if (newFile != null) {
                link.relink(newFile);
                link.update();
            }
        }
    }
    doc.save();
    doc.close();
        """ % (indd_file_js, dest_folder_js)

        try:
            # Uruchom ExtendScript. Drugi argument 1246973031 oznacza JavaScript/ExtendScript.
            result = indesign.DoScript(js, 1246973031)
            print("ExtendScript wykonany, wynik:", result)
        except Exception as e:
            messagebox.showerror("Błąd", "Błąd przy uruchamianiu skryptu ExtendScript:\n" + str(e))
            print(traceback.format_exc())

# ----------------------------------------
# CustomScrolledFrame – własna implementacja przewijalnego widgetu
# ----------------------------------------
class CustomScrolledFrame(tk.Frame):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, **kwargs)
        self.canvas = tk.Canvas(self, borderwidth=0, highlightthickness=0)
        self.vscrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=self.vscrollbar.set)
        self.vscrollbar.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.inner = tk.Frame(self.canvas)
        self.inner_id = self.canvas.create_window((0,0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", self.onFrameConfigure)
        self.canvas.bind("<Configure>", self.onCanvasConfigure)
        self.canvas.bind("<Enter>", self._bind_mousewheel)
        self.canvas.bind("<Leave>", self._unbind_mousewheel)
    
    def onFrameConfigure(self, event):
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))
    
    def onCanvasConfigure(self, event):
        self.canvas.itemconfig(self.inner_id, width=event.width)
    
    def _on_mousewheel(self, event):
        self.canvas.yview_scroll(int(-1*(event.delta/120)), "units")
    
    def _bind_mousewheel(self, event):
        self.canvas.bind_all("<MouseWheel>", self._on_mousewheel)
    
    def _unbind_mousewheel(self, event):
        self.canvas.unbind_all("<MouseWheel>")
    
    def display_widget(self, widget_class, **kwargs):
        widget = widget_class(self.inner, **kwargs)
        widget.pack(fill="both", expand=True)
        return widget

# ----------------------------------------
# DagaFrame – działanie modułu DAGA
# ----------------------------------------
class DagaFrame(tk.Frame):
    def __init__(self, master, *args, **kwargs):
        super().__init__(master, *args, **kwargs)
        wrapper = ttk.LabelFrame(self, text="Moduł DAGA (CSV -> InDesign)")
        wrapper.pack(fill="both", expand=True, padx=10, pady=10)
        self.scrolled_frame = CustomScrolledFrame(wrapper, width=500, height=500)
        self.scrolled_frame.pack(side="top", expand=True, fill="both", padx=5, pady=5)
        self.data_frame = self.scrolled_frame.inner
        
        self.data_groups = []
        
        self.name_entry_var = tk.StringVar()
        self.object_info_entry_var = tk.StringVar()
        self.object_info_pl_entry_var = tk.StringVar()
        self.size_entry_var = tk.StringVar()
        self.org_size_entry_var = tk.StringVar()
        self.quantity_color_size_entry_var = tk.StringVar()
        self.quantity_entry_var = tk.StringVar()
        self.image_org_sample_entry_var = tk.StringVar()
        self.author_entry_var = tk.StringVar()
        
        for var in (self.name_entry_var,
                    self.object_info_entry_var,
                    self.object_info_pl_entry_var,
                    self.size_entry_var,
                    self.org_size_entry_var,
                    self.quantity_color_size_entry_var,
                    self.quantity_entry_var,
                    self.image_org_sample_entry_var,
                    self.author_entry_var):
            var.trace_add("write", lambda *args, v=var: self.remove_quotes(v))
        
        self.create_basic_data_group()
        self.create_data_group()
        
        button_frame = tk.Frame(wrapper)
        button_frame.pack(fill="x", padx=5, pady=5)
        self.export_button = ttk.Button(button_frame, text="Eksportuj (CSV)", command=self.export_to_csv)
        self.export_button.grid(row=0, column=0, padx=5, pady=5)
        self.import_button = ttk.Button(button_frame, text="Importuj (CSV)", command=self.import_from_csv)
        self.import_button.grid(row=0, column=1, padx=5, pady=5)
        self.add_button = ttk.Button(button_frame, text="+", width=3, command=self.create_data_group)
        self.add_button.grid(row=0, column=2, padx=5, pady=5)
        self.remove_button = ttk.Button(button_frame, text="-", width=3, command=self.remove_data_group)
        self.remove_button.grid(row=0, column=3, padx=5, pady=5)
        self.reset_button = ttk.Button(button_frame, text="Resetuj", command=self.reset_to_defaults)
        self.reset_button.grid(row=0, column=4, padx=5, pady=5)
        self.help_button = ttk.Button(button_frame, text="Poradnik", command=self.open_help)
        self.help_button.grid(row=0, column=5, padx=5, pady=5)
        
        sv_ttk.set_theme("dark")
    
    def remove_quotes(self, var):
        old_value = var.get()
        new_value = old_value.replace('"', '')
        if new_value != old_value:
            var.set(new_value)

    def apply_data_merge(self, csv_file):
        """
        Ustawia źródło scalania danych (CSV) w wybranym pliku INDD.
        Operacje na łączach wykonuje funkcja w zakładce FOLDER.
        """
        try:
            import win32com.client
        except ImportError:
            messagebox.showerror("Błąd", "Moduł pywin32 nie jest zainstalowany.")
            return

        indd_file = filedialog.askopenfilename(
            title="Wybierz plik INDD",
            filetypes=[("InDesign Files", "*.indd")]
        )
        if not indd_file:
            return

        try:
            indesign = win32com.client.Dispatch("InDesign.Application")
            doc = indesign.Open(indd_file)
            doc.DataMergeProperties.SelectDataSource(csv_file)
            doc.Save()
            messagebox.showinfo("Sukces", "Scalanie danych zostało zastosowane.")
            doc.Close()
        except Exception as e:
            messagebox.showerror("Błąd", f"Błąd podczas scalania danych:\n{e}")

    def create_basic_data_group(self):
        frame = tk.Frame(self.data_frame)
        frame.pack(padx=10, pady=5, fill="x")
        
        tk.Label(frame, text="DATE:").grid(row=0, column=0, padx=10, pady=5, sticky="e")
        self.date_entry = DateEntry(frame, date_pattern='dd.mm.yyyy',
                                    year=date.today().year,
                                    month=date.today().month,
                                    day=date.today().day)
        self.date_entry.grid(row=0, column=1, padx=10, pady=5, sticky="w")

        tk.Label(frame, text="AUTHOR:").grid(row=1, column=0, padx=10, pady=5, sticky="e")
        self.author_entry = ttk.Entry(frame, textvariable=self.author_entry_var)
        self.author_entry.grid(row=1, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.author_entry)
            
        tk.Label(frame, text="NAME:").grid(row=2, column=0, padx=10, pady=5, sticky="e")
        self.name_entry = ttk.Entry(frame, textvariable=self.name_entry_var)
        self.name_entry.grid(row=2, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.name_entry)
        
        tk.Label(frame, text="OBJECTINFO:").grid(row=3, column=0, padx=10, pady=5, sticky="e")
        self.object_info_entry = ttk.Entry(frame, textvariable=self.object_info_entry_var)
        self.object_info_entry.grid(row=3, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.object_info_entry)
        
        tk.Label(frame, text="OBJINF_PL:").grid(row=4, column=0, padx=10, pady=5, sticky="e")
        self.object_info_pl_entry = ttk.Entry(frame, textvariable=self.object_info_pl_entry_var)
        self.object_info_pl_entry.grid(row=4, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.object_info_pl_entry)
        
        tk.Label(frame, text="SIZE:").grid(row=5, column=0, padx=10, pady=5, sticky="e")
        self.size_entry = ttk.Entry(frame, textvariable=self.size_entry_var)
        self.size_entry.grid(row=5, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.size_entry)
        
        tk.Label(frame, text="ORG_SIZE:").grid(row=6, column=0, padx=10, pady=5, sticky="e")
        self.org_size_entry = ttk.Entry(frame, textvariable=self.org_size_entry_var)
        self.org_size_entry.grid(row=6, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.org_size_entry)
        
        tk.Label(frame, text="QUANTITY_COLOR&SIZE:").grid(row=7, column=0, padx=10, pady=5, sticky="e")
        self.quantity_color_size_entry = ttk.Entry(frame, textvariable=self.quantity_color_size_entry_var)
        self.quantity_color_size_entry.grid(row=7, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.quantity_color_size_entry)
        
        tk.Label(frame, text="QUANTITY:").grid(row=8, column=0, padx=10, pady=5, sticky="e")
        self.quantity_entry = ttk.Entry(frame, textvariable=self.quantity_entry_var)
        self.quantity_entry.grid(row=8, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.quantity_entry)
        calculate_button = ttk.Button(frame, text="Wylicz", command=self.calculate_quantity)
        calculate_button.grid(row=8, column=2, padx=5, pady=5)
        
        tk.Label(frame, text="@IMAGE ORG_SAMPLE:").grid(row=9, column=0, padx=10, pady=5, sticky="e")
        self.image_org_sample_entry = ttk.Entry(frame, textvariable=self.image_org_sample_entry_var)
        self.image_org_sample_entry.grid(row=9, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(self.image_org_sample_entry)
        image_org_sample_button = ttk.Button(frame, text="Wybierz obraz",
                                             command=lambda: self.select_image(self.image_org_sample_entry))
        image_org_sample_button.grid(row=9, column=2, padx=5, pady=5)
 
    def create_data_group(self):
        frame = tk.Frame(self.data_frame)
        frame.pack(padx=10, pady=5, fill="x")
        group_number = len(self.data_groups) + 1
        color_entry_var = tk.StringVar()
        image_color_entry_var = tk.StringVar()
        codetxt_entry_var = tk.StringVar()
        image_codebr_entry_var = tk.StringVar()
        for var in (color_entry_var, image_color_entry_var, codetxt_entry_var, image_codebr_entry_var):
            var.trace_add("write", lambda *args, v=var: self.remove_quotes(v))
        tk.Label(frame, text=f"COLOR{group_number}:").grid(row=0, column=0, padx=10, pady=5, sticky="e")
        color_entry = ttk.Entry(frame, textvariable=color_entry_var)
        color_entry.grid(row=0, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(color_entry)
        
        tk.Label(frame, text=f"@IMAGE COLOR{group_number}:").grid(row=1, column=0, padx=10, pady=5, sticky="e")
        image_color_entry = ttk.Entry(frame, textvariable=image_color_entry_var)
        image_color_entry.grid(row=1, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(image_color_entry)
        image_color_button = ttk.Button(frame, text="Wybierz obraz",
                                        command=lambda: self.select_image(image_color_entry))
        image_color_button.grid(row=1, column=2, padx=5, pady=5)
        
        tk.Label(frame, text=f"CODETXT{group_number}:").grid(row=2, column=0, padx=10, pady=5, sticky="e")
        codetxt_entry = ttk.Entry(frame, textvariable=codetxt_entry_var)
        codetxt_entry.grid(row=2, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(codetxt_entry)
        
        tk.Label(frame, text=f"@IMAGE CODEBR{group_number}:").grid(row=3, column=0, padx=10, pady=5, sticky="e")
        image_codebr_entry = ttk.Entry(frame, textvariable=image_codebr_entry_var)
        image_codebr_entry.grid(row=3, column=1, padx=10, pady=5, sticky="w")
        add_entry_context_menu(image_codebr_entry)
        image_codebr_button = ttk.Button(frame, text="Wybierz obraz",
                                         command=lambda: self.select_image(image_codebr_entry))
        image_codebr_button.grid(row=3, column=2, padx=5, pady=5)
        self.data_groups.append((frame, color_entry_var, image_color_entry_var, codetxt_entry_var, image_codebr_entry_var))
    
    def select_image(self, entry):
        filename = filedialog.askopenfilename(
            title="Wybierz obraz",
            filetypes=[("Image Files", "*.png;*.jpg;*.jpeg;*.pdf;*.psd;*.ai")]
        )
        if filename:
            filename = filename.replace("/", "\\")
            entry.delete(0, tk.END)
            entry.insert(tk.END, filename)
    
    def calculate_quantity(self):
        size_text = self.size_entry_var.get().strip()
        sizes_list = [s.strip() for s in size_text.split(',') if s.strip()] if size_text else []
        size_count = len(sizes_list)
        color_count = sum(1 for group in self.data_groups if group[1].get().strip())
        try:
            qty_color_size = float(self.quantity_color_size_entry_var.get().strip())
        except ValueError:
            qty_color_size = 0.0
        total_qty = size_count * color_count * qty_color_size
        self.quantity_entry_var.set(str(int(total_qty)))
    
    def export_to_csv(self):
        mismatches = []
        for i, group in enumerate(self.data_groups, start=1):
            codetxt_value = group[3].get().strip()
            image_codebr_value = group[4].get().strip()
            if codetxt_value and image_codebr_value and (codetxt_value not in image_codebr_value):
                mismatches.append(f"Kolor {i}: {codetxt_value} nie pasuje do @IMAGE CODEBR")
        if mismatches:
            info_msg = ("Wykryto niezgodności w kodach EAN:\n\n" +
                        "\n".join(mismatches) +
                        "\n\nCzy kontynuować?")
            if not messagebox.askyesno("Niezgodność kodów EAN", info_msg):
                return
        object_info = self.object_info_entry_var.get()
        name = self.name_entry_var.get()
        default_filename = f"{object_info} {name}.csv"
        filename = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv")],
            initialfile=default_filename
        )
        if filename:
            with open(filename, 'w', newline='', encoding="utf-8") as file:
                writer = csv.writer(file, quoting=csv.QUOTE_MINIMAL)
                header = [
                    "DATE", "AUTHOR", "NAME", "OBJECTINFO", "OBJINF_PL", "SIZE", "ORG_SIZE",
                    "QUANTITY_COLOR&SIZE", "QUANTITY", "@IMAGE ORG_SAMPLE"
                ]
                for i in range(1, 7):
                    header.extend([f"COLOR{i}", f"@IMAGE COLOR{i}", f"CODETXT{i}", f"@IMAGE CODEBR{i}"])
                writer.writerow(header)
                # Poprawny porządek: najpierw DATE, potem AUTHOR, a następnie NAME itp.
                data_row = [
                    self.date_entry.get(),
                    self.author_entry_var.get(),  # Używamy pola AUTHOR jako drugi element
                    self.name_entry_var.get(),
                    self.object_info_entry_var.get(),
                    self.object_info_pl_entry_var.get(),
                    self.size_entry_var.get(),
                    self.org_size_entry_var.get(),
                    self.quantity_color_size_entry_var.get(),
                    self.quantity_entry_var.get(),
                    self.image_org_sample_entry_var.get()
                ]
                for group in self.data_groups[:6]:
                    data_row.extend([
                        group[1].get(),
                        group[2].get(),
                        group[3].get(),
                        group[4].get()
                    ])
                for _ in range(6 - len(self.data_groups)):
                    data_row.extend([""] * 4)
                writer.writerow(data_row)
            messagebox.showinfo("Eksport CSV", f"Plik CSV został zapisany:\n{filename}")
            
            if messagebox.askyesno("Scalanie danych", "Czy chcesz zastosować scalanie danych?"):
                self.apply_data_merge(filename)

    def import_from_csv(self):
        filename = filedialog.askopenfilename(filetypes=[("CSV files", "*.csv")])
        if filename:
            with open(filename, newline='', encoding="utf-8") as file:
                reader = csv.reader(file)
                header = next(reader)
                data = next(reader)
                self.date_entry.set_date(data[0])
                self.author_entry_var.set(data[1])
                self.name_entry_var.set(data[2])
                self.object_info_entry_var.set(data[3])
                self.object_info_pl_entry_var.set(data[4])
                self.size_entry_var.set(data[5])
                self.org_size_entry_var.set(data[6])
                self.quantity_color_size_entry_var.set(data[7])
                self.quantity_entry_var.set(data[8])
                self.image_org_sample_entry_var.set(data[9])
                for group in self.data_groups:
                    group[0].destroy()
                self.data_groups.clear()
                i = 10
                while i < len(data):
                    self.create_data_group()
                    group = self.data_groups[-1]
                    group[1].set(data[i])
                    group[2].set(data[i+1])
                    group[3].set(data[i+2])
                    group[4].set(data[i+3])
                    i += 4
        
    def reset_to_defaults(self):
        self.date_entry.set_date(date.today())
        self.author_entry_var.set("")
        self.name_entry_var.set("")
        self.object_info_entry_var.set("")
        self.object_info_pl_entry_var.set("")
        self.size_entry_var.set("")
        self.org_size_entry_var.set("")
        self.quantity_color_size_entry_var.set("")
        self.quantity_entry_var.set("")
        self.image_org_sample_entry_var.set("")
        for group in self.data_groups:
            group[0].destroy()
        self.data_groups.clear()
        self.create_data_group()
    
    def remove_data_group(self):
        if len(self.data_groups) > 1:
            last_group = self.data_groups.pop()
            last_group[0].destroy()
    
    def open_help(self):
        help_text = (
            "Wprowadzenie:\n"
            "D.A.G.A. (Dowolne Automatyczne Generowanie Artykułów) – generuje CSV do InDesign.\n\n"
            "Instrukcje:\n"
            "1. Wypełnij pola podstawowe.\n"
            "2. Dodaj grupy kolorów przy pomocy przycisków + / -.\n"
            "3. Kliknij 'Wylicz' aby obliczyć QUANTITY.\n"
            "4. Użyj Import/Export do pracy z plikiem CSV.\n\n"
            "Autor: Kamil Wróbel"
        )
        messagebox.showinfo("Poradnik", help_text)
    
    def fill_colors_from_list(self, color_list):
        for group in self.data_groups:
            group[0].destroy()
        self.data_groups.clear()
        for color in color_list:
            self.create_data_group()
            group = self.data_groups[-1]
            group[1].set(color)

# ----------------------------------------
# KeyNoteCoderFrame – pozostała funkcjonalność
# ----------------------------------------
class KeyNoteCoderFrame(tk.Frame):
    def __init__(self, master, daga_reference, *args, **kwargs):
        super().__init__(master, *args, **kwargs)
        self.daga_ref = daga_reference
        self.create_widgets()
        
    def create_widgets(self):
        top_label = ttk.Label(self, text=(
            "1) Utwórz notatnik z listą kolorów.\n"
            "2) Wygeneruj kody kreskowe."
        ))
        top_label.pack(pady=5)
        
        self.top_frame = ttk.Labelframe(self, text="Generowanie notatnika")
        self.top_frame.pack(fill="both", expand=True, padx=10, pady=5)
        self.bottom_frame = ttk.Labelframe(self, text="Generowanie kodów kreskowych")
        self.bottom_frame.pack(fill="both", expand=True, padx=10, pady=5)
        
        self.create_keynote_widgets(self.top_frame)
        self.create_coder_widgets(self.bottom_frame)
    
    def create_keynote_widgets(self, parent):
        prod_frame = tk.Frame(parent)
        prod_frame.pack(pady=5, fill="x")
        tk.Label(prod_frame, text="Nazwa produktu:").pack(side="left", padx=5)
        self.product_name_entry = ttk.Entry(prod_frame, width=50)
        self.product_name_entry.pack(side="left", padx=5)
        add_entry_context_menu(self.product_name_entry)
        
        colors_frame = tk.Frame(parent)
        colors_frame.pack(pady=5, fill="x")
        tk.Label(colors_frame, text="Kolory (1 linijka = 1 kolor):").pack(side="left", padx=5)
        self.colors_entry = tk.Text(colors_frame, height=5, width=30)
        self.colors_entry.pack(side="left", fill="x", expand=True, padx=5)
        add_text_context_menu(self.colors_entry)
        scroll = ttk.Scrollbar(colors_frame, command=self.colors_entry.yview)
        self.colors_entry.config(yscrollcommand=scroll.set)
        scroll.pack(side="left", fill="y")
        
        btn_frame = tk.Frame(parent)
        btn_frame.pack(pady=5)
        ttk.Button(btn_frame, text="Dodaj pustą linię", command=self.add_color_entry).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Usuń ostatnią linię", command=self.remove_color_entry).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Generuj notatnik (TXT)", command=self.create_notes).pack(side="left", padx=5)
        
        self.note_status_label = tk.Label(parent, text="", fg="cyan")
        self.note_status_label.pack(pady=5)
    
    def create_notes(self, *args):
        product_name = self.product_name_entry.get().strip()
        colors = [line.strip() for line in self.colors_entry.get("1.0", "end").splitlines() if line.strip()]
        if not product_name or not colors:
            self.note_status_label.config(text="Wprowadź nazwę produktu i co najmniej jedną linijkę kolorów.")
            return
        safe_product_name = product_name.replace('"', '')
        file_path = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Text files", "*.txt")],
            initialfile=f"{safe_product_name}.txt"
        )
        if file_path:
            with open(file_path, "w", encoding="utf-8") as file:
                max_len = max(len(product_name), *[len(c) for c in colors])
                for color in colors:
                    num_spaces = max_len - len(product_name)
                    line = f'{product_name}{" " * num_spaces} {color.upper()}{" " * (max_len - len(color))}\n'
                    file.write(line.upper())
                file.write("\n")
                excluded_words = ["MĘSKI", "MĘSKIE", "MĘSKA", "DAMSKI", "DAMSKIE", "DAMSKA"]
                for color in colors:
                    num_spaces = max_len - len(product_name)
                    line = f'{product_name}{" " * num_spaces} {color.upper()}{" " * (max_len - len(color))}\n'
                    line_without_excluded = " ".join(w for w in line.split() if w.upper() not in excluded_words)
                    file.write(line_without_excluded.upper() + "\n")
            self.note_status_label.config(text="Notatnik został utworzony i zapisany.")
        else:
            self.note_status_label.config(text="Nie wybrano pliku do zapisu.")
    
    def add_color_entry(self):
        self.colors_entry.insert("end", "\n")
    
    def remove_color_entry(self):
        all_text = self.colors_entry.get("1.0", "end")
        lines = all_text.splitlines()
        if lines:
            self.colors_entry.delete(f"{len(lines)}.0", "end")
    
    def create_coder_widgets(self, parent):
        file_frame = tk.Frame(parent)
        file_frame.pack(pady=5, fill="x")
        ttk.Button(file_frame, text="Generuj z notatnika (TXT)", command=self.generate_from_file).pack(side="left", padx=5)
        format_frame = ttk.LabelFrame(file_frame, text="Format wyjściowy")
        format_frame.pack(side="left", padx=10)
        self.output_format_var = tk.StringVar(value="jpg")
        ttk.Radiobutton(format_frame, text="JPG", variable=self.output_format_var, value="jpg").pack(side="left", padx=5)
        ttk.Radiobutton(format_frame, text="PDF", variable=self.output_format_var, value="pdf").pack(side="left", padx=5)
        text_frame = tk.Frame(parent)
        text_frame.pack(pady=5, fill="x")
        self.text_area = tk.Text(text_frame, height=5)
        self.text_area.pack(side="left", fill="x", expand=True)
        add_text_context_menu(self.text_area)
        ttk.Button(text_frame, text="Generuj z pola tekstowego", command=self.generate_from_text).pack(side="right", padx=5)
    
    def generate_ean13_images(self, input_file, output_dir, output_format="jpg"):
        with open(input_file, 'r', encoding="utf-8") as file:
            lines = file.readlines()
        for line in lines:
            line = line.strip()
            match = re.search(r'\b\d{13}\b', line)
            if match:
                code = match.group()
                url = (f"https://bwipjs-api.metafloor.com/?bcid=ean13"
                       f"&text={code}"
                       f"&scale=3"
                       f"&height=90"
                       f"&includetext"
                       f"&textsize=12"
                       f"&backgroundcolor=ffffff"
                       f"&barcolor=000000")
                response = requests.get(url)
                if response.status_code == 200:
                    try:
                        image = Image.open(BytesIO(response.content))
                        if image.mode == 'RGBA':
                            image = image.convert('RGB')
                        original_width, original_height = image.size
                        crop_height = 138
                        cropped_image = image.crop((0, original_height - crop_height, original_width, original_height))
                        margin_width = 22
                        final_width = cropped_image.width + margin_width
                        final_image = Image.new('RGB', (final_width, cropped_image.height), 'white')
                        final_image.paste(cropped_image, (0, 0))
                        filename_base = os.path.join(output_dir, f"ean-13_{code}")
                        if output_format.lower() == "jpg":
                            final_image.save(f"{filename_base}.jpg", 'JPEG')
                        elif output_format.lower() == "pdf":
                            final_image.save(f"{filename_base}.pdf", 'PDF')
                    except Exception as e:
                        print(f"Błąd przetwarzania kodu {code}: {e}")
    
    def generate_from_file(self):
        input_file = filedialog.askopenfilename(title="Wybierz plik z kodami EAN-13",
                                                filetypes=[("Text files", "*.txt")])
        if input_file:
            output_dir = os.path.dirname(input_file)
            fmt = self.output_format_var.get()
            self.generate_ean13_images(input_file, output_dir, output_format=fmt)
            messagebox.showinfo("Sukces", f"Kody zostały wygenerowane i zapisane jako {fmt.upper()}!")
    
    def generate_from_text(self):
        codes = self.text_area.get("1.0", "end").strip().split("\n")
        extracted_codes = []
        for code_line in codes:
            extracted_codes.extend(re.findall(r'\b\d{13}\b', code_line))
        if not extracted_codes:
            messagebox.showerror("Błąd", "Brak poprawnych kodów EAN-13 w polu tekstowym.")
            return
        output_dir = filedialog.askdirectory(title="Wybierz folder wyjściowy")
        if output_dir:
            temp_file = os.path.join(output_dir, "temp_codes.txt")
            with open(temp_file, 'w', encoding="utf-8") as f:
                f.write("\n".join(extracted_codes))
            fmt = self.output_format_var.get()
            self.generate_ean13_images(temp_file, output_dir, output_format=fmt)
            os.remove(temp_file)
            messagebox.showinfo("Sukces", f"Kody z pola tekstowego zostały wygenerowane jako {fmt.upper()}!")
    
# ----------------------------------------
# Główne uruchomienie aplikacji
# ----------------------------------------
class MainApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("DAGA X")
        self.iconbitmap(r"S:\Graficy Specyfikacje\Techniczne\_PROGRAMY\program files\bin\DAGAX.ico")
        
        self.bind("<Control-q>", lambda e: self.quit())
        self.bind("<Control-n>", lambda e: self.knc_tab.create_notes())
        self.bind("<Control-b>", lambda e: self.knc_tab.generate_from_file())
        self.bind("<Control-e>", lambda e: self.daga_tab.export_to_csv())
        self.bind("<Control-i>", lambda e: self.daga_tab.import_from_csv())
        self.bind("<Control-r>", lambda e: self.daga_tab.reset_to_defaults())
        
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True)
        
        self.folder_creation_tab = FolderCreationFrame(self.notebook)
        self.notebook.add(self.folder_creation_tab, text="FOLDER")
        
        self.knc_tab = KeyNoteCoderFrame(self.notebook, daga_reference=None)
        self.notebook.add(self.knc_tab, text="CODE")
        
        self.daga_tab = DagaFrame(self.notebook)
        self.notebook.add(self.daga_tab, text="DAGA")
        
        self.excel_tab = ExcelTab(self.notebook)
        try:
            img = Image.open(EXCEL_ICON_PATH)
            if img.mode != "RGBA":
                img = img.convert("RGBA")
            r, g, b, a = img.split()
            rgb_image = Image.merge("RGB", (r, g, b))
            inverted_rgb = ImageOps.invert(rgb_image)
            inverted_image = Image.merge("RGBA", (*inverted_rgb.split(), a))
            resized_icon = inverted_image.resize((18, 18), Image.LANCZOS)
            self.excel_icon = ImageTk.PhotoImage(resized_icon)
        except Exception as e:
            messagebox.showerror("Błąd", f"Nie udało się wczytać ikony Excel:\n{EXCEL_ICON_PATH}\n\n{e}")
            self.excel_icon = None
        self.notebook.add(self.excel_tab, text="PROJEKTY", image=self.excel_icon, compound="left")
        
        sv_ttk.set_theme("dark")

def main():
    app = MainApp()
    app.mainloop()

if __name__ == "__main__":
    main()
